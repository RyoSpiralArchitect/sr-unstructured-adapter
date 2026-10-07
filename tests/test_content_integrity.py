# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regressions for source content preservation and parser dispatch."""

import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from sr_adapter import pipeline
from sr_adapter.embedding import EmbeddingIndex
from sr_adapter.loaders import read_file_contents
from sr_adapter.parsers import parse_csv, parse_html, parse_json, parse_pptx, parse_txt, parse_xlsx
from sr_adapter.profiles import ProcessingProfile, ProfileStore, resolve_profile
from sr_adapter.recipe import _build_recipe_config, load_recipe
from sr_adapter.refiner import HybridRefiner
from sr_adapter.schema import BBox, Block, DocumentMeta, Span
from sr_adapter.sniff import detect_type


def test_sniff_pdf_magic_without_suffix_uses_bounded_read(tmp_path, monkeypatch):
    source = tmp_path / "document.bin"
    source.write_bytes(b"%PDF-1.7\n" + b"x" * 100)
    monkeypatch.setattr(Path, "read_bytes", lambda self: pytest.fail("must not read entire file"))
    assert detect_type(source) == "pdf"


def test_xlsx_preserves_values_and_all_sheets(tmp_path):
    source = tmp_path / "table.xlsx"
    book = Workbook()
    book.active.title = "First"
    book.active.append(["name", 0, False])
    book.create_sheet("Second").append(["猫", 42])
    book.save(source)
    book.close()
    blocks = parse_xlsx(source)
    assert [block.attrs["sheet"] for block in blocks] == ["First", "Second"]
    assert json.loads(blocks[0].attrs["rows"]) == [["name", "0", "False"]]
    assert json.loads(blocks[1].attrs["rows"]) == [["猫", "42"]]
    doc = pipeline.convert(source, "default", llm_ok=False)
    assert [block.text for block in doc.blocks] == [block.text for block in blocks]
    assert not doc.warnings


def test_xlsx_loader_obeys_exact_row_limit(tmp_path, monkeypatch):
    source = tmp_path / "table.xlsx"
    book = Workbook()
    book.active.append(["first"])
    book.active.append(["second"])
    book.save(source)
    book.close()
    monkeypatch.setenv("SR_ADAPTER_XLSX_MAX_ROWS", "1")
    text, meta = read_file_contents(source, "application/octet-stream")
    assert text == "first"
    assert meta["xlsx_rows_read"] == 1


def test_tsv_preserves_quoted_cells_and_bom(tmp_path):
    source = tmp_path / "table.tsv"
    source.write_text('key\tvalue\nname\t"a\tb"\n', encoding="utf-8-sig")
    assert json.loads(parse_csv(source)[0].attrs["rows"]) == [["key", "value"], ["name", "a\tb"]]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32", "cp932"])
def test_text_parser_preserves_encoded_japanese(tmp_path, encoding):
    source = tmp_path / "message.txt"
    text = "これは日本語の文章です。"
    source.write_text(text, encoding=encoding)
    assert parse_txt(source)[0].text == text


def test_json_repairs_do_not_rewrite_quoted_content(tmp_path):
    source = tmp_path / "config.json"
    source.write_text('{"url": "https://example.org/a", "literal": "/*keep*/,}", // comment\n "ok": true,}', encoding="utf-8")
    blocks = parse_json(source)
    values = {block.attrs.get("key"): block.attrs.get("value") for block in blocks}
    assert values["url"] == "https://example.org/a"
    assert values["literal"] == "/*keep*/,}"
    assert values["ok"] == "true"


def test_html_preserves_code_once_and_drops_scripts(tmp_path):
    source = tmp_path / "sample.html"
    source.write_text('<script>secret()</script><p>Use <code>API_KEY</code>.</p><pre><code>if True:\n    print("OK")\n</code></pre>', encoding="utf-8")
    blocks = parse_html(source)
    assert len(blocks) == 2
    assert blocks[1].type == "code"
    assert blocks[1].text == 'if True:\n    print("OK")\n'
    assert all("secret" not in block.text for block in blocks)


def test_pptx_slides_sort_numerically(tmp_path):
    source = tmp_path / "slides.pptx"
    with zipfile.ZipFile(source, "w") as archive:
        for number in [10, 2, 1]:
            archive.writestr(f"ppt/slides/slide{number}.xml", f'<slide xmlns:a="urn:a"><a:t>Slide {number}</a:t></slide>')
    assert [block.text for block in parse_pptx(source)] == ["Slide 1", "Slide 2", "Slide 10"]


def test_registered_pdf_override_is_used_for_streaming(tmp_path, monkeypatch):
    source = tmp_path / "input.pdf"
    source.write_bytes(b"%PDF-custom")
    monkeypatch.setitem(pipeline.REGISTRY.by_key, "pdf", lambda path: [Block(text="custom parser")])
    assert pipeline._parse(source, detected="pdf", mime=None)[0].text == "custom parser"
    assert list(pipeline._stream_raw(source, detected="pdf", mime=None))[0].text == "custom parser"


def test_partial_stream_failure_is_not_replaced_with_file_text(tmp_path, monkeypatch):
    source = tmp_path / "input.pdf"
    source.write_bytes(b"%PDF-binary")
    def stream(path):
        yield Block(text="page one")
        raise ValueError("page two failed")
    monkeypatch.setitem(pipeline._STREAMERS, "pdf", stream)
    result = pipeline._stream_raw(source, detected="pdf", mime=None)
    assert next(result).text == "page one"
    with pytest.raises(ValueError, match="page two failed"):
        next(result)


def test_binary_parse_failure_is_observable_without_decoding_bytes(tmp_path):
    source = tmp_path / "broken.xlsx"
    source.write_bytes(b"PK\x03\x04sensitive-binary")
    document = pipeline.convert(source, "default", llm_ok=False)
    assert document.warnings == ["xlsx parser failed (BadZipFile)"]
    assert document.blocks[0].text == ""


def test_stream_honors_profile_block_limit(tmp_path):
    source = tmp_path / "input.md"
    source.write_text("one paragraph\nsecond paragraph\nthird paragraph", encoding="utf-8")
    profile = ProcessingProfile(name="limited", max_blocks=1)
    assert len(list(pipeline.stream_convert(source, "default", profile=profile))) == 1


def test_refiner_preserves_case_code_spans_and_confidence():
    refiner = HybridRefiner()
    blocks = [
        Block(text="  API NASA  ", confidence=0.1, attrs={"ml_refine": "unexpected"}),
        Block(type="code", text="if TRUE:\n    CALL()\n", confidence=0.1),
        Block(text="  label  ", spans=[Span(start=2, end=7)], confidence=0.1),
    ]
    results = refiner.refine(blocks)
    assert results[0].text == "API NASA"
    assert results[1].text == blocks[1].text
    assert results[2].text == blocks[2].text
    assert results[2].spans == blocks[2].spans
    assert all(block.confidence == 0.1 for block in results)


@pytest.mark.parametrize("name", ["../default", "/tmp/secret", "nested/name", "..\\default"])
def test_config_names_do_not_escape_resource_directories(name):
    with pytest.raises(ValueError):
        load_recipe(name)
    with pytest.raises(ValueError):
        ProfileStore().load(name)


def test_recipe_rejects_unknown_block_type():
    with pytest.raises(ValueError, match="Invalid recipe"):
        _build_recipe_config("bad", {"fallback": {"as": "invented-type"}})


def test_auto_profile_falls_back_when_selector_disabled(monkeypatch):
    monkeypatch.setattr("sr_adapter.profiles._maybe_auto_selector", lambda: None)
    assert resolve_profile("auto").name == "balanced"


def test_document_keeps_primary_language():
    assert DocumentMeta(languages=["ja"], primary_language="ja").model_dump()["primary_language"] == "ja"


@pytest.mark.parametrize("value", [float("nan"), float("inf")])
def test_bbox_rejects_nonfinite_coordinates(value):
    with pytest.raises(ValidationError):
        BBox(x0=value, y0=0, x1=1, y1=1)


def test_embedding_index_validates_search_and_finite_vectors():
    with pytest.raises(ValueError):
        EmbeddingIndex(0)
    index = EmbeddingIndex(2)
    with pytest.raises(ValueError, match="finite"):
        index.add([float("nan"), 0])
    index.add([1, 0])
    assert index.search([1, 0], top_k=0) == []
    with pytest.raises(ValueError, match="top_k"):
        index.search([1, 0], top_k=-1)
    with pytest.raises(ValueError, match="finite"):
        index.search([float("inf"), 0])
    with pytest.raises(ValueError):
        index.extend([[1, 0]], [])


def test_docx_keeps_interleaved_table_order(tmp_path):
    from docx import Document
    from sr_adapter.parsers import parse_docx

    source = tmp_path / "document.docx"
    document = Document()
    document.add_paragraph("Before the table.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table cell"
    document.add_paragraph("After the table.")
    document.save(source)
    assert [block.text for block in parse_docx(source)] == ["Before the table.", "Table cell", "After the table."]


def test_pptx_uses_presentation_order(tmp_path):
    source = tmp_path / "reordered.pptx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("ppt/presentation.xml", '<p:presentation xmlns:p="urn:p" xmlns:r="urn:r"><p:sldIdLst><p:sldId r:id="r2"/><p:sldId r:id="r1"/></p:sldIdLst></p:presentation>')
        archive.writestr("ppt/_rels/presentation.xml.rels", '<Relationships><Relationship Id="r1" Target="slides/slide1.xml"/><Relationship Id="r2" Target="slides/slide2.xml"/></Relationships>')
        for number in [1, 2]:
            archive.writestr(f"ppt/slides/slide{number}.xml", f'<slide xmlns:a="urn:a"><a:t>Slide {number}</a:t></slide>')
    blocks = parse_pptx(source)
    assert [block.text for block in blocks] == ["Slide 2", "Slide 1"]
    assert [block.prov.page for block in blocks] == [0, 1]


def test_calendar_preserves_parameters_folding_and_event_description(tmp_path):
    from sr_adapter.parsers import parse_ics

    source = tmp_path / "schedule.ics"
    source.write_text(
        "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Planning\\, team\n"
        "DESCRIPTION:First line\\nsecond\n  folded word\n"
        "DTSTART;TZID=Asia/Tokyo:20261008T090000\n"
        "BEGIN:VALARM\nDESCRIPTION:Alarm text\nEND:VALARM\n"
        "END:VEVENT\nEND:VCALENDAR\n", encoding="utf-8",
    )
    block = parse_ics(source)[0]
    assert block.attrs["summary"] == "Planning, team"
    assert block.attrs["description"] == "First line\nsecond folded word"
    assert block.attrs["dtstart"] == "20261008T090000"
    assert block.attrs["dtstart_params"] == "TZID=Asia/Tokyo"
    assert "Start: 20261008T090000" in block.text
    assert "Alarm text" not in block.text


def test_zero_escalation_threshold_is_honored():
    from sr_adapter.escalation.model import LinearEscalationModel
    from sr_adapter.escalation.policy import EscalationPolicyEngine
    from sr_adapter.settings import EscalationSettings

    engine = EscalationPolicyEngine(
        model=LinearEscalationModel({}, bias=-1.0, threshold=0.5),
        settings=EscalationSettings(min_score=0.0),
    )
    assert engine.threshold == 0.0
    assert engine.select([Block(text="Review this", confidence=0.1)]) == [0]


def test_image_exif_diagnostics_are_json_serializable(tmp_path, monkeypatch):
    import sys
    from PIL import Image

    monkeypatch.setitem(sys.modules, "pytesseract", None)
    source = tmp_path / "metadata.jpg"
    exif = Image.Exif()
    exif[37510] = b"ASCII\x00\x00\x00comment"
    Image.new("RGB", (8, 8)).save(source, exif=exif)
    _, metadata = read_file_contents(source, "image/jpeg")
    assert "UserComment" in metadata["image_exif_sample"]
    json.dumps(metadata)


@pytest.mark.parametrize("native_disabled", [False, True])
def test_pdf_text_needs_no_native_layout_and_keeps_estimates_out_of_provenance(monkeypatch, tmp_path, native_disabled):
    from types import SimpleNamespace
    from sr_adapter.parsers import stream_pdf

    if native_disabled:
        monkeypatch.setenv("SR_ADAPTER_DISABLE_NATIVE_RUNTIME", "on")
    else:
        monkeypatch.delenv("SR_ADAPTER_DISABLE_NATIVE_RUNTIME", raising=False)
    text = "this paragraph contains ordinary extracted text with no real source coordinates."
    monkeypatch.setattr("sr_adapter.parsers.PdfReader", lambda path: SimpleNamespace(pages=[
        SimpleNamespace(extract_text=lambda: text),
        SimpleNamespace(extract_text=lambda: "a paragraph on the second page."),
    ]))
    monkeypatch.setattr("sr_adapter.parsers.VisualLayoutAnalyzer", lambda **kwargs: pytest.fail("PDF text must not require native layout"))
    blocks = list(stream_pdf(tmp_path / "document.pdf"))
    assert [block.prov.page for block in blocks] == [0, 1]
    assert [block.prov.order for block in blocks] == [0, 0]
    assert all(block.prov.bbox is None for block in blocks)
    assert all(block.type == "paragraph" and block.confidence == 0.5 for block in blocks)
    assert all(block.attrs["layout_geometry"] == "heuristic" for block in blocks)
    assert all("layout_bbox_estimate" in block.attrs for block in blocks)
    assert all("layout_confidence" not in block.attrs for block in blocks)


@pytest.mark.parametrize("native_disabled", [False, True])
def test_image_text_and_real_ocr_provenance_survive_unavailable_layout(monkeypatch, tmp_path, native_disabled):
    from sr_adapter.parsers import stream_image

    if native_disabled:
        monkeypatch.setenv("SR_ADAPTER_DISABLE_NATIVE_RUNTIME", "true")
    else:
        monkeypatch.delenv("SR_ADAPTER_DISABLE_NATIVE_RUNTIME", raising=False)
    monkeypatch.setattr("sr_adapter.parsers._extract_image_text", lambda path: (
        "OCR text", {"image_has_text": True}, [{
            "text": "OCR text", "source": "ocr", "kind": "ocr", "confidence": 0.25,
            "bbox": (2.0, 3.0, 42.0, 53.0), "page": 2, "order": 7,
        }],
    ))
    def missing_layout(**kwargs):
        if native_disabled:
            pytest.fail("disabled native layout must not be constructed")
        raise RuntimeError("compiler not found")
    monkeypatch.setattr("sr_adapter.parsers.VisualLayoutAnalyzer", missing_layout)
    blocks = list(stream_image(tmp_path / "image.png"))
    block = next(block for block in blocks if block.attrs.get("image_source") == "ocr")
    assert block.text == "OCR text"
    assert block.confidence == 0.25
    assert block.type == "paragraph"
    assert block.prov.page == 2 and block.prov.order == 7
    assert block.prov.bbox == BBox(x0=2, y0=3, x1=42, y1=53)
    assert block.attrs["layout_geometry"] == "ocr"
    assert block.attrs["layout_analysis"] == ("disabled" if native_disabled else "unavailable")


def test_image_synthetic_geometry_never_drives_classification(monkeypatch, tmp_path):
    from sr_adapter.parsers import stream_image

    monkeypatch.delenv("SR_ADAPTER_DISABLE_NATIVE_RUNTIME", raising=False)
    monkeypatch.setattr("sr_adapter.parsers._extract_image_text", lambda path: (
        "text without coordinates", {}, [{
            "text": "text without coordinates", "source": "ocr", "kind": "ocr",
            "confidence": 0.3, "page": 1, "order": 4,
        }],
    ))
    monkeypatch.setattr("sr_adapter.parsers.VisualLayoutAnalyzer", lambda **kwargs: pytest.fail("heuristic geometry must not drive layout"))
    block = list(stream_image(tmp_path / "image.png"))[0]
    assert block.text == "text without coordinates"
    assert block.type == "paragraph" and block.confidence == 0.3
    assert block.prov.bbox is None
    assert block.prov.page == 1 and block.prov.order == 4
    assert block.attrs["layout_geometry"] == "heuristic"
    assert block.attrs["layout_analysis"] == "skipped_heuristic_geometry"


def test_image_layout_failure_does_not_duplicate_partial_output(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from sr_adapter.parsers import stream_image

    monkeypatch.delenv("SR_ADAPTER_DISABLE_NATIVE_RUNTIME", raising=False)
    segments = [{"text": f"line {i}", "source": "ocr", "confidence": 0.4,
                 "bbox": (0, i * 20, 100, i * 20 + 10), "order": i}
                for i in range(2)]
    monkeypatch.setattr("sr_adapter.parsers._extract_image_text", lambda path: ("", {}, segments))
    class FailingAnalyzer:
        def __init__(self, **kwargs):
            pass
        def process(self, candidates):
            yield SimpleNamespace(block=candidates[0].block)
            raise RuntimeError("kernel failure")
    monkeypatch.setattr("sr_adapter.parsers.VisualLayoutAnalyzer", FailingAnalyzer)
    blocks = list(stream_image(tmp_path / "image.png"))
    assert [block.text for block in blocks] == ["line 0", "line 1"]
    assert all(block.attrs["layout_analysis"] == "unavailable" for block in blocks)
