"""Generate fixed CSV/PPTX/PNG inputs from memo-a.md, without reading gold.

Requires Pillow and python-pptx. Pass a local Japanese-capable font explicitly;
the font is used for rendering but is never copied into the fixture directory.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
from importlib.metadata import version
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo


HERE = Path(__file__).resolve().parent
SECTION = "## 担当と予算"


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def write_pptx(path: Path, lines: list[str], assertions: list[tuple[int, str]], font_name: str) -> list[dict]:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = lines[0]
    presentation.core_properties.author = "SR synthetic fixture"
    fixed_time = datetime(2026, 10, 17, tzinfo=timezone.utc)
    presentation.core_properties.created = fixed_time
    presentation.core_properties.modified = fixed_time
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    def textbox(text: str, name: str, top: float, height: float, size: int) -> int:
        shape = slide.shapes.add_textbox(Inches(0.4), Inches(top), Inches(12.5), Inches(height))
        shape.name = name
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_top = frame.margin_bottom = 0
        frame.text = text
        for paragraph in frame.paragraphs:
            paragraph.font.name = font_name
            paragraph.font.size = Pt(size)
            paragraph.space_after = 0
        return len(slide.shapes) - 1

    textbox(lines[0], "document-title", 0.25, 0.45, 23)
    context = "\n".join(line for line in lines[1:lines.index(SECTION) + 1] if line)
    textbox(context, "document-scope", 0.85, 2.0, 14)
    locations = []
    for index, (source_line, text) in enumerate(assertions):
        shape_index = textbox(text, f"assertion-{index + 1:02d}", 3.0 + index * 0.5, 0.45, 18)
        locations.append({"source_line": source_line, "slide_number": 1, "shape_index": shape_index,
                          "shape_name": f"assertion-{index + 1:02d}", "text": text})
    buffer = BytesIO()
    presentation.save(buffer)
    # Freeze ZIP entry ordering/timestamps; the fixed assets are the evaluation input.
    with ZipFile(buffer) as original, ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=9) as frozen:
        for name in sorted(original.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            frozen.writestr(info, original.read(name), compress_type=ZIP_DEFLATED, compresslevel=9)
    return locations


def wrap_line(text: str, font, width: int) -> list[str]:
    if not text:
        return [""]
    rows = []
    current = ""
    for character in text:
        candidate = current + character
        if current and font.getlength(candidate) > width:
            rows.append(current)
            current = character
        else:
            current = candidate
    rows.append(current)
    return rows


def write_png(path: Path, lines: list[str], font_path: Path, font_index: int) -> tuple[list[dict], dict]:
    from PIL import Image, ImageDraw, ImageFont

    width, margin, font_size, line_height = 1800, 70, 32, 48
    font = ImageFont.truetype(str(font_path), font_size, index=font_index)
    rows = [(number, wrap_line(text, font, width - 2 * margin)) for number, text in enumerate(lines, 1)]
    height = 2 * margin + sum(len(wrapped) * line_height for _, wrapped in rows)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    positions = []
    y = margin
    for source_line, wrapped in rows:
        start = y
        for line in wrapped:
            draw.text((margin, y), line, fill="black", font=font)
            y += line_height
        if lines[source_line - 1]:
            positions.append({"source_line": source_line, "page_number": 1,
                              "bbox_pixels": [margin, start, width - margin, y],
                              "text": lines[source_line - 1]})
    image.save(path, optimize=True)
    return positions, {"width": width, "height": height, "font_size_pixels": font_size,
                       "line_height_pixels": line_height, "font_family": list(font.getname())}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=HERE / "memo-a.md")
    parser.add_argument("--output-dir", type=Path, default=HERE / "variants")
    parser.add_argument("--font", type=Path, required=True, help="Local Japanese-capable TTF/OTF/TTC; not redistributed.")
    parser.add_argument("--font-index", type=int, default=0, help="Font face index for a TTC collection.")
    parser.add_argument("--overwrite", action="store_true", help="Explicitly replace the fixed generated assets.")
    args = parser.parse_args()
    source = args.source.resolve()
    raw = source.read_bytes()
    lines = raw.decode("utf-8").splitlines()
    if lines.count(SECTION) != 1:
        parser.error("Source must contain exactly one 担当と予算 section.")
    assertions = [(number, line) for number, line in enumerate(lines, 1)
                  if number > lines.index(SECTION) + 1 and line.strip()]
    if len(assertions) != 8 or any("の" not in text for _, text in assertions):
        parser.error("This fixed variant requires eight subject-bearing statements in source order.")
    if not args.font.is_file() or args.font_index < 0:
        parser.error("Provide an existing Japanese-capable font and a nonnegative face index.")
    from PIL import ImageFont

    font_name = ImageFont.truetype(str(args.font), 32, index=args.font_index).getname()[0]
    versions = {package: version(package) for package in ("Pillow", "python-pptx")}
    paths = {kind: args.output_dir / f"{source.stem}.{kind}" for kind in ("csv", "pptx", "png")}
    manifest_path = args.output_dir / "manifest.json"
    if not args.overwrite and any(path.exists() for path in [*paths.values(), manifest_path]):
        parser.error("Fixed outputs already exist; choose another output directory or use --overwrite.")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    assertion_numbers = {number for number, _ in assertions}
    csv_locations = []
    with paths["csv"].open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["案件", "記述"])
        row_number = 1
        for source_line, text in enumerate(lines, 1):
            if not text:
                continue
            row_number += 1
            subject = text.split("の", 1)[0] if source_line in assertion_numbers else "資料情報"
            writer.writerow([subject, text])
            if source_line in assertion_numbers:
                csv_locations.append({"source_line": source_line, "row_number": row_number,
                                      "column_number": 2, "column_name": "記述", "text": text})
    pptx_locations = write_pptx(paths["pptx"], lines, assertions, font_name)
    png_lines, rendering = write_png(paths["png"], lines, args.font, args.font_index)
    png_locations = [line for line in png_lines if line["source_line"] in assertion_numbers]
    manifest = {
        "schema_version": 1,
        "source_document": source.name,
        "source_sha256": digest(raw),
        "assertion_count": len(assertions),
        "gold_read": False,
        "independent_document_count": 0,
        "metadata_preserved": "Document title and identity/precedence scope are included in every format.",
        "generator_sha256": digest(Path(__file__).read_bytes()),
        "libraries": versions,
        "font": {"file_name": args.font.name, "sha256": digest(args.font.read_bytes()),
                 "face_index": args.font_index, "font_file_redistributed": False},
        "rendering": rendering,
        "formats": {},
    }
    for kind, locations in [("csv", csv_locations), ("pptx", pptx_locations), ("png", png_locations)]:
        manifest["formats"][kind] = {"file": paths[kind].name, "sha256": digest(paths[kind].read_bytes()),
                                     "bytes": paths[kind].stat().st_size, "assertion_locations": locations}
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"manifest": str(manifest_path), "formats": list(paths), "assertions_per_format": 8},
                     ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
